#!/usr/bin/env python3
"""
"상품바코드조회" 형태의 엑셀 파일을 앱이 오프라인으로 읽는
assets/product_barcodes.json 으로 변환하는 스크립트.

상품 마스터 데이터가 갱신될 때마다 이 스크립트를 다시 실행해서
assets/product_barcodes.json을 새로 만들고, 앱을 다시 빌드하면 됩니다.
(인터넷 연결 없이 기기에 데이터가 그대로 들어가는 방식이라, 데이터가
바뀔 때는 파일을 새로 만들어 앱을 재배포해야 합니다.)

사용법:
    pip install openpyxl
    python3 convert_barcode_xlsx_to_json.py 상품바코드조회.xlsx

엑셀 열 순서가 조금씩 바뀌어도(예: 맨 앞에 "바코드구분" 열이 추가되거나,
"바코드" 열 이름이 "상품코드"로 바뀌는 등) 안정적으로 동작하도록, 열 번호를
고정하지 않고 "헤더 이름"으로 열을 찾습니다. 헤더 줄도 1행이 아니어도(예:
1행이 "기본사항/세부사항" 같은 그룹 제목이고 2행이 실제 헤더인 경우) 자동으로
찾습니다.
"""

import json
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl이 필요합니다. 먼저 `pip install openpyxl`을 실행하세요.")

# 컬럼 데이터 키 -> 엑셀 헤더에 쓰일 수 있는 이름(들). 여러 개를 적어두면
# 그중 먼저 발견되는 이름을 사용합니다(엑셀 양식이 바뀌어도 대응하기 위함).
COLUMN_NAME_CANDIDATES = {
    "itemNo": ["품번"],
    "name": ["품명"],
    "color": ["색상"],
    "size": ["사이즈"],
    "sizeLabel": ["사이즈표기"],
    "barcode": ["바코드", "상품코드"],
    "category": ["사이즈구분명"],
    "gender": ["성별구분명"],
    "price": ["최초판매가"],
    "tagPrice": ["택가"],
}

REQUIRED_KEYS = ["itemNo", "barcode"]


def _find_header_row(ws, max_scan_rows: int = 5):
    """맨 위 몇 줄 중에서 실제 컬럼 헤더가 있는 행을 찾습니다.
    ("품번"과 바코드 열 이름 중 하나가 함께 있는 첫 행)"""
    barcode_names = set(COLUMN_NAME_CANDIDATES["barcode"])
    for r in range(1, max_scan_rows + 1):
        values = [str(c.value).strip() if c.value is not None else "" for c in ws[r]]
        if "품번" in values and any(name in values for name in barcode_names):
            return r, values
    raise SystemExit(
        "엑셀에서 헤더 행을 찾지 못했습니다. 위쪽 5개 행 안에 '품번'과 "
        "'바코드'(또는 '상품코드') 열 이름이 함께 있는지 확인해주세요."
    )


def _build_column_index(header_values):
    index = {}
    missing = []
    for key, candidates in COLUMN_NAME_CANDIDATES.items():
        found_col = None
        for name in candidates:
            if name in header_values:
                found_col = header_values.index(name) + 1  # openpyxl은 1부터 시작
                break
        if found_col is None:
            if key in REQUIRED_KEYS:
                missing.append(key)
        else:
            index[key] = found_col
    if missing:
        raise SystemExit(f"필수 컬럼을 찾지 못했습니다: {missing}")
    return index


def convert(xlsx_path: str, out_path: str, sheet_name: str | None = None) -> int:
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb[sheet_name] if sheet_name else wb[wb.sheetnames[0]]

    header_row, header_values = _find_header_row(ws)
    column_index = _build_column_index(header_values)

    def clean_text(v) -> str:
        # 엑셀/XML 변환 과정에서 가끔 섞여 들어오는 "_x000D_" 같은
        # 캐리지리턴 잔재를 제거합니다.
        s = str(v)
        return s.replace("_x000D_", "").strip()

    records = []
    for r in range(header_row + 1, ws.max_row + 1):
        def cell(col_key):
            col = column_index.get(col_key)
            return ws.cell(row=r, column=col).value if col else None

        item_no = cell("itemNo")
        barcode = cell("barcode")
        if not item_no or not barcode:
            continue  # 품번/바코드가 없는 빈 행은 건너뜀

        price = cell("price")
        tag_price = cell("tagPrice")

        records.append({
            "itemNo": clean_text(item_no),
            "name": clean_text(cell("name") or ""),
            "color": clean_text(cell("color") or "-"),
            "size": clean_text(cell("size") or ""),
            "sizeLabel": clean_text(cell("sizeLabel") or ""),
            "barcode": clean_text(barcode),
            "category": clean_text(cell("category") or ""),
            "gender": clean_text(cell("gender") or ""),
            "price": int(price) if isinstance(price, (int, float)) else None,
            "tagPrice": int(tag_price) if isinstance(tag_price, (int, float)) else None,
        })

    Path(out_path).write_text(
        json.dumps(records, ensure_ascii=False, separators=(",", ":")),
        encoding="utf-8",
    )
    return len(records)


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit("사용법: python3 convert_barcode_xlsx_to_json.py <엑셀파일.xlsx>")

    input_path = sys.argv[1]
    output_path = str(Path(__file__).resolve().parent.parent / "assets" / "product_barcodes.json")

    count = convert(input_path, output_path)
    print(f"완료: {count}건 -> {output_path}")
